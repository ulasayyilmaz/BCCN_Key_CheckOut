import os
import json
import threading
from datetime import datetime
from flask import Flask, request, jsonify
from flask_cors import CORS
import openpyxl

# STEP 3 ADDED: Imports for PDF generation
import base64
import io
import traceback
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader

app = Flask(__name__)
CORS(app)

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
EXCEL_FILE = os.path.join(BASE_DIR, "checkouts.xlsx")
HEADER = ["Index", "Name", "Student ID", "Key ID", "Purpose", "Checkout Time", "Check-in Time", "PDF Filename"]

# Global lock for file operations and index assignment
file_lock = threading.Lock()

def get_next_index_and_write(row_data):
    """
    writes row_data (given in argument) to the excel sheet with the right index

    """
    with file_lock:
        if not os.path.exists(EXCEL_FILE):
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(HEADER)
            next_index = 1
        else:
            try:
                # loads the spreadsheet
                workbook = openpyxl.load_workbook(EXCEL_FILE)
                sheet = workbook.active

                # Determine next index from the last row
                if sheet.max_row <= 1:
                    next_index = 1
                else:
                    # Read the last index (last row, first column)
                    last_index_val = sheet.cell(row=sheet.max_row, column=1).value
                    try:
                        next_index = int(last_index_val) + 1
                    except (ValueError, TypeError):
                        # Fallback if the cell value isn't a valid integer
                        next_index = sheet.max_row
            except Exception as e:
                raise Exception(f"Failed to load Excel file: {e}")
        
        # Prepend the index to the row_data
        full_row = [next_index] + row_data
        sheet.append(full_row)
        
        try:
            workbook.save(EXCEL_FILE)
        except Exception as e:
            raise Exception(f"Failed to save Excel file: {e}")
            
        return next_index

# STEP 3 ADDED: PDF Generation function
def generate_pdf(index, data, checkout_time):
    try:
        pdf_dir = os.path.join(BASE_DIR, "pdfs")
        if not os.path.exists(pdf_dir):
            os.makedirs(pdf_dir)
            
        safe_name = data['name'].replace(" ", "_")
        pdf_filename = f"{index}_{safe_name}.pdf"
        pdf_path = os.path.join(pdf_dir, pdf_filename)
        
        c = canvas.Canvas(pdf_path, pagesize=letter)
        width, height = letter
        
        # Header
        c.setFont("Helvetica-Bold", 16)
        c.drawString(50, height - 50, "BCCN Berlin — Key Checkout Form")
        c.setLineWidth(0.5)
        c.line(50, height - 60, width - 50, height - 60)
        
        # Fields
        c.setFont("Helvetica", 12)
        c.drawString(50, height - 100, f"Record #: {index}")
        c.drawString(50, height - 120, f"Name: {data['name']}")
        c.drawString(50, height - 140, f"Student ID: {data['student_id']}")
        c.drawString(50, height - 160, f"Key ID: {data['key_id']}")
        c.drawString(50, height - 180, f"Purpose: {data['purpose']}")
        c.drawString(50, height - 200, f"Checkout Time: {checkout_time}")
        
        # Signature
        c.drawString(50, height - 240, "Signature:")
        
        # Decode signature
        img_data = base64.b64decode(data['signature_b64'])
        img_io = io.BytesIO(img_data)
        img = ImageReader(img_io)
        c.drawImage(img, 50, height - 400, width=300, height=150, mask='auto')
        
        # Footer
        c.setFont("Helvetica", 10)
        c.setFillColorRGB(0.5, 0.5, 0.5)
        c.drawString(50, 50, f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M:%S')} — Record {index}")
        
        c.save()
        return pdf_filename
    except Exception as e:
        print(f"Error generating PDF: {e}")
        traceback.print_exc()
        return None

# flask is listening to local server,
# whenever http request posted, it receives the data 
# by request.get_json
@app.route('/submit', methods=['POST'])
def submit():
    try:
        # gets data by get_json (posted via ngrok tunnel from netlify server)
        data = request.get_json()
        if not data:
            return jsonify({"success": False, "error": "Invalid or missing JSON payload"}), 400
            
        # checks for required fields
        required_fields = ["name", "student_id", "key_id", "purpose", "signature_b64"]
        
        for field in required_fields:
            if field not in data or not isinstance(data[field], str) or not data[field].strip():
                return jsonify({"success": False, "error": f"Missing or empty required field: '{field}'"}), 400
                
        # Record checkout time (server's local time)
        checkout_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # Prepare row data from 
        row_data = [
            data["name"].strip(),
            data["student_id"].strip(),
            data["key_id"].strip(),
            data["purpose"].strip(),
            checkout_time,
            "",  # Check-in Time
            ""   # PDF Filename
        ]
        
        print("Writing to Excel...")
        # gets index and writes to excel file
        next_index = get_next_index_and_write(row_data)
        print("Done. Row index:", next_index)
        
        # STEP 3 ADDED: Generate PDF and update Excel
        pdf_filename = generate_pdf(next_index, data, checkout_time)
        if pdf_filename:
            try:
                with file_lock:
                    workbook = openpyxl.load_workbook(EXCEL_FILE)
                    sheet = workbook.active
                    # Find the row and update PDF Filename (column 8)
                    for row in range(sheet.max_row, 1, -1):
                        if sheet.cell(row=row, column=1).value == next_index:
                            sheet.cell(row=row, column=8).value = pdf_filename
                            break
                    workbook.save(EXCEL_FILE)
            except Exception as e:
                print(f"Error updating Excel with PDF filename: {e}")
        
        # returns json of the data with index
        return jsonify({"success": True, "index": next_index}), 200
        
    except Exception as e:
        return jsonify({"success": False, "error": str(e)}), 500

if __name__ == '__main__':
    app.run(debug=False, use_reloader=False, port=5000)

